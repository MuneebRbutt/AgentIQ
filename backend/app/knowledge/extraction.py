"""Turn an uploaded file into plain text (FR-05-01). Supported: .txt, .md, .pdf, .docx."""

import io
from pathlib import PurePath

from app.core.errors import ValidationError

SUPPORTED_EXTENSIONS = (".txt", ".md", ".pdf", ".docx")


class UnsupportedDocumentError(ValidationError):
    code = "unsupported_document"


def extract_text(filename: str, data: bytes) -> str:
    """The document's text, or `UnsupportedDocumentError` if it can't be read.

    Blocking (PDF parsing can take a while): call it from a worker thread.
    """
    extension = PurePath(filename).suffix.lower()
    try:
        if extension in (".txt", ".md"):
            text = _decode(data)
        elif extension == ".pdf":
            text = _pdf_text(data)
        elif extension == ".docx":
            text = _docx_text(data)
        else:
            raise UnsupportedDocumentError(
                f"Unsupported file type. Upload one of: {', '.join(SUPPORTED_EXTENSIONS)}."
            )
    except UnsupportedDocumentError:
        raise
    except Exception as exc:  # corrupt or encrypted files raise many different errors
        raise UnsupportedDocumentError(
            f"Couldn't read {filename}: the file seems damaged."
        ) from exc

    text = text.replace("\r\n", "\n").replace("\x00", "").strip()
    if not text:
        raise UnsupportedDocumentError(
            f"{filename} has no readable text (scanned PDFs without a text layer aren't supported)."
        )
    return text


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1")


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader  # imported lazily: only needed for PDFs

    reader = PdfReader(io.BytesIO(data))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _docx_text(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))
    return "\n\n".join(paragraphs)
