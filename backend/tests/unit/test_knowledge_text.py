"""Text extraction and chunking of uploaded documents (FR-05-01/02)."""

import io
import itertools

import pytest

from app.knowledge.chunking import chunk_text
from app.knowledge.extraction import UnsupportedDocumentError, extract_text


def _docx(*paragraphs: str, table: list[list[str]] | None = None) -> bytes:
    from docx import Document

    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _pdf(text: str) -> bytes:
    """A one-page PDF with one line of text, built by hand (no PDF writer needed)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(
        b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
        % (
            len(objects) + 1,
            xref,
        )
    )
    return out.getvalue()


class TestExtraction:
    def test_plain_text_in_common_encodings(self) -> None:
        assert extract_text("a.txt", b"Check-in at 2 PM.") == "Check-in at 2 PM."
        assert extract_text("a.txt", "﻿No pets.".encode()) == "No pets."
        assert extract_text("a.TXT", "Café open".encode("utf-16")) == "Café open"
        assert extract_text("a.md", "Prix: 5 €".encode("cp1252")).startswith("Prix: 5")

    def test_markdown(self) -> None:
        assert "Breakfast" in extract_text("faq.md", b"# FAQ\n\nBreakfast is at 7.")

    def test_docx_paragraphs_and_tables(self) -> None:
        data = _docx("Opening hours", table=[["Monday", "9-5"], ["Tuesday", "9-1"]])
        text = extract_text("hours.docx", data)
        assert "Opening hours" in text
        assert "Monday | 9-5" in text

    def test_pdf(self) -> None:
        assert "Cancel 24 hours ahead" in extract_text("policy.pdf", _pdf("Cancel 24 hours ahead"))

    def test_unsupported_type_is_rejected(self) -> None:
        with pytest.raises(UnsupportedDocumentError) as error:
            extract_text("photo.png", b"\x89PNG")
        assert error.value.code == "unsupported_document"

    def test_corrupt_file_is_rejected(self) -> None:
        with pytest.raises(UnsupportedDocumentError):
            extract_text("broken.pdf", b"not a pdf")
        with pytest.raises(UnsupportedDocumentError):
            extract_text("broken.docx", b"not a zip")

    def test_document_without_text_is_rejected(self) -> None:
        with pytest.raises(UnsupportedDocumentError):
            extract_text("empty.txt", b"  \n\n ")


class TestChunking:
    def test_short_text_is_one_chunk(self) -> None:
        assert chunk_text("Breakfast is at 7.\n\nNo pets.") == ["Breakfast is at 7.\n\nNo pets."]

    def test_chunks_respect_the_size_limit_and_keep_all_words(self) -> None:
        paragraphs = [f"Rule {n} says guests must do thing number {n} politely." for n in range(60)]
        chunks = chunk_text("\n\n".join(paragraphs), max_chars=200, overlap_chars=60)
        assert len(chunks) > 1
        assert all(len(chunk) <= 200 for chunk in chunks)
        joined = " ".join(chunks)
        assert all(f"thing number {n} " in joined for n in range(60))

    def test_sentences_are_not_cut_and_overlap_carries_context(self) -> None:
        sentences = [f"Sentence {n} is about policy {n}." for n in range(30)]
        chunks = chunk_text(" ".join(sentences), max_chars=150, overlap_chars=60)
        for chunk in chunks:
            assert chunk.endswith(".")
        # The last sentence of each chunk is repeated at the start of the next.
        for previous, following in itertools.pairwise(chunks):
            last_sentence = "Sentence" + previous.rsplit("Sentence", 1)[1]
            assert last_sentence in following

    def test_overlong_words_are_wrapped(self) -> None:
        chunks = chunk_text("x" * 450, max_chars=200, overlap_chars=0)
        assert [len(chunk) for chunk in chunks] == [200, 200, 50]

    def test_whitespace_is_normalised(self) -> None:
        assert chunk_text("Open   daily\nfrom 9.\n\n\n\nClosed  Sundays.") == [
            "Open daily from 9.\n\nClosed Sundays."
        ]

    def test_overlap_must_be_smaller_than_size(self) -> None:
        with pytest.raises(ValueError, match="overlap"):
            chunk_text("text", max_chars=100, overlap_chars=100)
